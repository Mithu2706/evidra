"""PPTX parsing with python-pptx: slide text, tables, speaker notes, links,
media and structural visibility signals. Rendering is delegated to a converter
(LibreOffice) that produces a PDF."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from ..config import get_settings
from .base import DocumentParser, IngestionError, ParsedDocument, ParsedUnit, TextSpanSignal
from .textutil import clean_text, contrast_ratio

TINY_FONT_PT = 4.0
LOW_CONTRAST_RATIO = 1.35
_P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _rgb_tuple(rgb) -> tuple[float, float, float]:
    s = str(rgb)
    return (int(s[0:2], 16) / 255, int(s[2:4], 16) / 255, int(s[4:6], 16) / 255)


def _solid_fill_rgb(fill) -> tuple[float, float, float] | None:
    try:
        if fill.type == MSO_FILL.SOLID and fill.fore_color.type is not None:
            return _rgb_tuple(fill.fore_color.rgb)
    except Exception:
        return None
    return None


class PPTXParser(DocumentParser):
    file_types = ("pptx",)
    name = "python-pptx"

    def parse(self, path: Path, work_dir: Path) -> ParsedDocument:
        try:
            prs = Presentation(str(path))
        except Exception as exc:
            raise IngestionError(
                "The presentation could not be opened; the file appears damaged or incomplete."
            ) from exc

        slide_w, slide_h = int(prs.slide_width or Emu(12192000)), int(prs.slide_height or Emu(6858000))
        units: list[ParsedUnit] = []
        for index, slide in enumerate(prs.slides):
            unit = ParsedUnit(page_number=index + 1)
            try:
                self._parse_slide(slide, unit, slide_w, slide_h)
            except Exception as exc:
                unit.processing_notes.append(f"Text extraction failed on this slide: {exc}")
            units.append(unit)
        if not units:
            raise IngestionError("Presentation has no slides.")

        notes: list[str] = []
        pdf_path = convert_to_pdf(path, work_dir)
        page_map: dict[int, int] = {}
        if pdf_path is None:
            notes.append("Slide rendering unavailable (LibreOffice not found or conversion failed).")
        else:
            import pymupdf as fitz

            with fitz.open(pdf_path) as pdf:
                n_pages = pdf.page_count
            visible = [i for i, u in enumerate(units) if not u.hidden]
            if n_pages == len(units):
                page_map = {i: i for i in range(len(units))}
            elif n_pages == len(visible):
                page_map = {unit_idx: page for page, unit_idx in enumerate(visible)}
            else:
                notes.append(
                    f"Rendered page count ({n_pages}) does not match slide count ({len(units)}); "
                    "slide images were not attached."
                )
                pdf_path = None

        return ParsedDocument(
            file_type="pptx",
            units=units,
            renderable_pdf=pdf_path,
            render_page_map=page_map,
            notes=notes,
        )

    # ------------------------------------------------------------------

    def _parse_slide(self, slide, unit: ParsedUnit, slide_w: int, slide_h: int) -> None:
        unit.hidden = slide._element.get("show") in ("0", "false")
        slide_bg = _solid_fill_rgb(slide.background.fill) if slide.follow_master_background is False else None
        background = slide_bg or (1.0, 1.0, 1.0)

        lines: list[str] = []
        links: set[str] = set()

        title_shape = slide.shapes.title
        if title_shape is not None and title_shape.has_text_frame:
            unit.title = clean_text(title_shape.text_frame.text).split("\n")[0][:200] or None

        for shape in self._iter_shapes(slide.shapes):
            st = shape.shape_type
            if st in (MSO_SHAPE_TYPE.PICTURE,) or getattr(shape, "has_chart", False):
                unit.has_visual_content = True
            if st in (MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
                unit.has_embedded_media = True
            try:
                if shape.click_action.hyperlink.address:
                    links.add(shape.click_action.hyperlink.address)
            except Exception:
                pass

            off_slide = self._is_off_slide(shape, slide_w, slide_h)
            hidden_shape = self._is_hidden_shape(shape)
            shape_fill = None
            try:
                shape_fill = _solid_fill_rgb(shape.fill)
            except Exception:
                pass

            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [clean_text(c.text) for c in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
                continue

            if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                para_text = "".join(r.text for r in para.runs).strip()
                if not para_text:
                    continue
                lines.append(para_text)
                for run in para.runs:
                    if run.hyperlink and run.hyperlink.address:
                        links.add(run.hyperlink.address)
                if off_slide:
                    unit.structural_signals.append(
                        TextSpanSignal(para_text, "off_slide_content", "Text box is positioned outside the visible slide.")
                    )
                    continue
                if hidden_shape:
                    unit.structural_signals.append(
                        TextSpanSignal(para_text, "invisible_text_layer", "Shape is marked hidden.")
                    )
                    continue
                reason = self._run_visibility(para.runs, shape_fill or background)
                if reason:
                    unit.structural_signals.append(TextSpanSignal(para_text, *reason))

        unit.text = clean_text("\n".join(lines))
        unit.links = sorted(links)
        if slide.has_notes_slide:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
            unit.speaker_notes = notes or None
        if unit.hidden:
            unit.processing_notes.append("Slide is marked hidden and is not shown in presentation mode.")

    @staticmethod
    def _iter_shapes(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from PPTXParser._iter_shapes(shape.shapes)
            else:
                yield shape

    @staticmethod
    def _is_off_slide(shape, slide_w: int, slide_h: int) -> bool:
        if shape.left is None or shape.top is None:
            return False
        w, h = shape.width or 0, shape.height or 0
        return shape.left >= slide_w or shape.top >= slide_h or shape.left + w <= 0 or shape.top + h <= 0

    @staticmethod
    def _is_hidden_shape(shape) -> bool:
        for el in shape._element.iter(f"{_P_NS}cNvPr"):
            return el.get("hidden") in ("1", "true")
        return False

    @staticmethod
    def _run_visibility(runs, background) -> tuple[str, str] | None:
        for run in runs:
            if not run.text.strip():
                continue
            size = run.font.size
            if size is not None and size.pt < TINY_FONT_PT:
                return ("tiny_text", f"Font size {size.pt:.1f}pt.")
            try:
                color = run.font.color
                if color and color.type is not None and color.rgb is not None:
                    if contrast_ratio(_rgb_tuple(color.rgb), background) < LOW_CONTRAST_RATIO:
                        return ("low_contrast_text", "Text colour is nearly identical to the background.")
            except (AttributeError, ValueError, TypeError):
                continue
        return None


def _find_libreoffice() -> str | None:
    configured = get_settings().libreoffice_path
    if configured:
        return configured if Path(configured).exists() else shutil.which(configured)
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def convert_to_pdf(path: Path, out_dir: Path, timeout: int = 180) -> Path | None:
    binary = _find_libreoffice()
    if not binary:
        return None
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="evidra-lo-") as profile:
        cmd = [
            binary,
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--norestore",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(path),
        ]
        try:
            subprocess.run(cmd, capture_output=True, timeout=timeout, check=False)
        except (subprocess.TimeoutExpired, OSError):
            return None
    pdf = out_dir / (path.stem + ".pdf")
    return pdf if pdf.exists() and pdf.stat().st_size > 0 else None
